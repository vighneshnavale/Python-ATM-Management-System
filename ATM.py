
import openpyxl
from openpyxl import Workbook
import os

FILE_NAME = "atm_data.xlsx"

def create_excel():
    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Accounts"
        ws.append(["Name", "Phone", "Balance"])
        wb.save(FILE_NAME)

def get_sheet():
    wb = openpyxl.load_workbook(FILE_NAME)
    return wb, wb["Accounts"]

def find_account(ws, phone):
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 2).value) == phone:
            return row
    return None

def create_account():
    wb, ws = get_sheet()

    name = input("Enter your name: ").strip()
    phone = input("Enter your phone number: ").strip()

    if not name or not phone.isdigit():
        print("Invalid name or phone number.")
        wb.close()
        return

    if find_account(ws, phone):
        print("Account already exists!")
        wb.close()
        return

    try:
        balance = float(input("Enter opening balance: "))
        if balance < 0:
            print("Balance cannot be negative.")
            wb.close()
            return
    except ValueError:
        print("Enter a valid amount.")
        wb.close()
        return

    ws.append([name, phone, balance])
    wb.save(FILE_NAME)
    wb.close()

    print("Account created successfully!")
    print("Your account has been created.")


def deposit_money():
    wb, ws = get_sheet()
    phone = input("Enter your phone number: ").strip()
    row = find_account(ws, phone)

    if row is None:
        print("Account not found!")
        wb.close()
        return

    try:
        amount = float(input("Enter deposit amount: "))
        if amount <= 0:
            print("Amount must be greater than zero.")
            wb.close()
            return
    except ValueError:
        print("Enter a valid amount.")
        wb.close()
        return

    balance = float(ws.cell(row, 3).value)
    ws.cell(row, 3, balance + amount)
    wb.save(FILE_NAME)
    wb.close()

    print("Money deposited successfully!")
    print("Updated balance:", balance + amount)


def withdraw_money():
    wb, ws = get_sheet()
    phone = input("Enter your phone number: ").strip()
    row = find_account(ws, phone)

    if row is None:
        print("Account not found!")
        wb.close()
        return

    try:
        amount = float(input("Enter withdrawal amount: "))
        if amount <= 0:
            print("Amount must be greater than zero.")
            wb.close()
            return
    except ValueError:
        print("Enter a valid amount.")
        wb.close()
        return

    balance = float(ws.cell(row, 3).value)

    if amount > balance:
        print("Insufficient balance!")
        print("Available balance:", balance)
        wb.close()
        return

    new_balance = balance - amount
    ws.cell(row, 3, new_balance)
    wb.save(FILE_NAME)
    wb.close()

    print("Please collect your money.")
    print("Withdrawal successful!")
    print("Remaining balance:", new_balance)


def check_balance():
    wb, ws = get_sheet()
    phone = input("Enter your phone number: ").strip()
    row = find_account(ws, phone)

    if row is None:
        print("Account not found!")
    else:
        print("\n----- ACCOUNT INFORMATION -----")
        print("Name:", ws.cell(row, 1).value)
        print("Phone:", ws.cell(row, 2).value)
        print("Current balance:", ws.cell(row, 3).value)
        print("--------------------------------")

    wb.close()


def main():
    create_excel()

    while True:
        print("\n========== ATM SYSTEM ==========")
        print("1. Create Account")
        print("2. Withdraw Money")
        print("3. Add Money")
        print("4. Check Balance")
        print("5. Exit")

        choice = input("Enter your choice: ").strip()

        if choice == "1":
            create_account()
        elif choice == "2":
            withdraw_money()
        elif choice == "3":
            deposit_money()
        elif choice == "4":
            check_balance()
        elif choice == "5":
            print("Thank you for using our ATM system!")
            break
        else:
            print("Invalid choice! Please select 1 to 5.")

if __name__ == "__main__":
    main()
